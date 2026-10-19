"""独立复算：直接从原始 XLSX 重算关键计数，不调用 survey-report / schema。

用途：验证 `docs/_work/survey/report.md` 的关键数字，而不是只验证脚本能跑。
规则（与导入脚本一致，但独立实现）：
- 多选题按分号拆分；“无”选项与具体选项同时出现时，剔除“无”选项；
- 多选题分母 30；单选题分母 30；
- 单选题必须检查全部选项之和等于有效回答数（Q10 的七类合计必须为 30），
  并单独核对“任何付费形式”“不打算个人付费”“四分类（不含条件性）”“五分类（含条件性）”；
- 时长 = 完成时间 − 开始时间（秒）。

运行：
  python docs/_work/survey/analysis/recompute_from_xlsx.py --input "<XLSX>"
输出：
  docs/_work/survey/analysis/recompute.md
"""

from __future__ import annotations

import argparse
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

NONE_OPTIONS = {
    8: {"以上皆无"},
    9: {"以上均没有"},
    10: {"没有明显影响", "没有遇到相关问题"},
    11: {"目前的方法已经足够好"},
    13: {"以上均没有明显吸引力"},
    14: {"没有明显担忧"},
}
MULTI_COLUMNS = {8: "Q3", 9: "Q4", 10: "Q5", 11: "Q6", 13: "Q8", 14: "Q9"}


def split_multi(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"[;；]", value) if item.strip()]


def count_option(rows: list[list[str]], column: int, option: str) -> int:
    total = 0
    for row in rows:
        selected = split_multi(row[column])
        none = NONE_OPTIONS.get(column, set())
        if any(item in none for item in selected) and any(item not in none for item in selected):
            selected = [item for item in selected if item not in none]
        total += 1 if option in selected else 0
    return total


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    args = parser.parse_args()

    workbook = load_workbook(Path(args.input), read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = [["" if cell is None else str(cell).strip() for cell in row] for row in list(sheet.iter_rows(values_only=True))[1:]]
    if len(rows) != 30:
        raise SystemExit(f"预期 30 行，实际 {len(rows)}")

    durations = []
    for row in rows:
        start = datetime.strptime(row[1], "%Y-%m-%d %H:%M:%S")
        end = datetime.strptime(row[2], "%Y-%m-%d %H:%M:%S")
        durations.append(int((end - start).total_seconds()))

    q7 = Counter(row[12] for row in rows)
    q10 = Counter(row[15] for row in rows)
    q1 = Counter(row[6] for row in rows)
    any_paid = q10["为一次完整的消费计划支付少量费用"] + q10["为持续使用支付月费"] + q10["为高级功能支付一次性费用"]
    no_personal_pay = q10["只使用免费版本，不愿付费"] + q10["由银行、学校或其他机构提供免费服务"]
    seven_sum = sum(q10.values())
    four_way = any_paid + q10["只使用免费版本，不愿付费"] + q10["由银行、学校或其他机构提供免费服务"] + q10["暂时无法判断"]
    five_way = four_way + q10["视实际功能和效果而定"]

    checks = [
        ("Q1 18-20 岁", q1.get("18-20", 0), 22),
        ("Q3 跨城观演", count_option(rows, 8, "跨城演唱会、音乐节或现场演出"), 9),
        ("Q3 旅行", count_option(rows, 8, "旅行"), 22),
        ("Q3 电子产品", count_option(rows, 8, "购买较昂贵的电子产品或其他商品"), 19),
        ("Q3 以上皆无", count_option(rows, 8, "以上皆无"), 4),
        ("Q4 至少 1 项问题", 30 - count_option(rows, 9, "以上均没有"), 22),
        ("Q4 担心影响生活", count_option(rows, 9, "担心消费影响日常生活开支"), 11),
        ("Q4 付款日不足", count_option(rows, 9, "总体资金足够，但某个付款日期前可用资金不足"), 4),
        ("Q4 超预算", count_option(rows, 9, "总费用超出预算"), 9),
        ("Q4 忘记/低估后续费用", count_option(rows, 9, "忘记或低估后续费用"), 8),
        ("Q4 推迟或放弃", count_option(rows, 9, "因不确定是否负担得起而推迟或放弃计划"), 6),
        ("Q4 难比较总成本", count_option(rows, 9, "难以比较不同方案的总成本"), 6),
        ("Q4 以上均没有", count_option(rows, 9, "以上均没有"), 8),
        ("Q5 至少 1 项影响", 30 - count_option(rows, 10, "没有明显影响") - count_option(rows, 10, "没有遇到相关问题"), 22),
        ("Q5 焦虑或犹豫", count_option(rows, 10, "感到焦虑或犹豫"), 17),
        ("Q5 减少其他开支", count_option(rows, 10, "减少其他开支"), 11),
        ("Q5 动用储备", count_option(rows, 10, "动用原本希望保留的储备资金"), 8),
        ("Q5 额外时间", count_option(rows, 10, "花费额外时间计算和核对"), 7),
        ("Q5 临时修改或取消", count_option(rows, 10, "临时修改或取消计划"), 6),
        ("Q6 至少 1 项不足", 30 - count_option(rows, 11, "目前的方法已经足够好"), 26),
        ("Q6 计划变化需重算", count_option(rows, 11, "计划变化时需要重新计算"), 13),
        ("Q6 手工输入", count_option(rows, 11, "需要手动收集和输入太多资料"), 9),
        ("Q6 难比较完整方案", count_option(rows, 11, "难以比较多个完整消费方案"), 8),
        ("Q6 现有方法已够好", count_option(rows, 11, "目前的方法已经足够好"), 4),
        ("Q6 没有固定规划方式", count_option(rows, 11, "没有固定的规划方式"), 7),
        ("Q6 难掌握付款日期", count_option(rows, 11, "难以掌握每笔费用的付款日期"), 6),
        ("Q7 一定会尝试", q7.get("一定会尝试", 0), 8),
        ("Q7 很可能会尝试", q7.get("很可能会尝试", 0), 13),
        ("Q7 不确定", q7.get("不确定", 0), 7),
        ("Q7 不太可能", q7.get("不太可能", 0), 1),
        ("Q7 一定不会", q7.get("一定不会", 0), 1),
        ("Q7 试用前两档", q7.get("一定会尝试", 0) + q7.get("很可能会尝试", 0), 21),
        ("Q8 多方案比较", count_option(rows, 13, "多方案比较： 比较不同门票、交通、住宿组合的费用与付款安排"), 13),
        ("Q8 付款时间轴", count_option(rows, 13, "付款时间轴： 查看每笔费用的付款日期及对应资金需求"), 11),
        ("Q8 必要资金保护", count_option(rows, 13, "必要资金保护： 保留日常生活费及指定储备资金"), 9),
        ("Q8 资金缺口预警", count_option(rows, 13, "资金缺口预警： 提前发现某个时间点可能出现的资金不足"), 9),
        ("Q8 情景模拟", count_option(rows, 13, "情景模拟： 测试收入延迟或费用增加后的影响"), 2),
        ("Q8 不确定费用提示", count_option(rows, 13, "不确定费用提示： 标记尚未确认的费用及其影响"), 2),
        ("Q8 无吸引", count_option(rows, 13, "以上均没有明显吸引力"), 4),
        ("Q9 隐私", count_option(rows, 14, "个人财务信息的隐私与安全"), 16),
        ("Q9 遗漏费用/过于乐观", count_option(rows, 14, "是否遗漏费用或给出过于乐观的判断"), 9),
        ("Q9 输入麻烦", count_option(rows, 14, "输入资料是否太麻烦"), 9),
        ("Q9 计算准确性", count_option(rows, 14, "计算结果是否准确"), 6),
        ("Q9 易理解", count_option(rows, 14, "结果是否容易理解"), 3),
        ("Q9 无明显担忧", count_option(rows, 14, "没有明显担忧"), 2),
        ("Q10 只用免费版", q10["只使用免费版本，不愿付费"], 12),
        ("Q10 机构免费", q10["由银行、学校或其他机构提供免费服务"], 5),
        ("Q10 按次付费", q10["为一次完整的消费计划支付少量费用"], 5),
        ("Q10 月费", q10["为持续使用支付月费"], 2),
        ("Q10 一次性付费", q10["为高级功能支付一次性费用"], 1),
        ("Q10 视功能与效果而定", q10["视实际功能和效果而定"], 1),
        ("Q10 暂时无法判断", q10["暂时无法判断"], 4),
        ("Q10 七类合计（单选必须等于 30）", seven_sum, 30),
        ("Q10 四分类合计（不含条件性）", four_way, 29),
        ("Q10 五分类合计（含条件性）", five_way, 30),
        ("Q10 任何付费形式", any_paid, 8),
        ("Q10 不打算个人付费", no_personal_pay, 17),
        ("时长最小值（秒）", min(durations), 21),
        ("时长中位数（秒）", sorted(durations)[14:16], [84, 92]),
        ("时长最大值（秒）", max(durations), 277),
        ("不足 60 秒", sum(1 for value in durations if value < 60), 8),
    ]

    lines = [
        "# 问卷关键数字独立复算（不调用 survey-report）",
        "",
        "> 输入：原始 XLSX（SHA-256 见 `report.md`）；本脚本独立实现分号拆分、\"无\"选项修正与计数，",
        "> 只用于核对 `report.md`，不替代统计表。",
        "",
        "| 检查项 | 独立复算 | report.md | 一致 |",
        "| --- | --- | --- | --- |",
    ]
    failed = 0
    for name, actual, expected in checks:
        same = actual == expected
        failed += 0 if same else 1
        lines.append(f"| {name} | {actual} | {expected} | {'是' if same else '否'} |")
    lines.append("")
    lines.append(f"结论：{len(checks)} 项中 {len(checks) - failed} 项一致，{failed} 项不一致。")
    lines.append("")
    out = Path(__file__).parent / "recompute.md"
    out.write_text("\n".join(lines), encoding="utf-8", newline="\n")
    print(f"已写入 {out}：{len(checks) - failed}/{len(checks)} 一致")
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
